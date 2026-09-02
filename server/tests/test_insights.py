"""Forecasting, anomalies, budgets, break-even, benchmark, unit economics,
menu items, month-close pack."""
import io
import openpyxl


def _today():
    from datetime import date
    return date.today().isoformat()


def test_forecast_shape_and_math(client, outlet_id):
    client.post("/api/auth/stepup", json={"password": "change-me-please"})
    # use LAST month (fully past) so the math is deterministic
    from datetime import date, timedelta
    t = date.today()
    first_this = t.replace(day=1)
    hi_prev = first_this - timedelta(days=1)
    ym = f"{hi_prev.year}-{hi_prev.month:02d}"
    dim = (hi_prev - hi_prev.replace(day=1)).days + 1

    d = hi_prev.replace(day=1)
    while d <= hi_prev:
        rr = client.put("/api/sales/manual", headers=None, json={
            "outlet_id": outlet_id, "business_date": d.isoformat(),
            "channel_kind": "cash", "amount_rupees": 1000})
        assert rr.status_code == 200, f"seed {d} → {rr.status_code} {rr.text[:120]}"
        d += timedelta(days=1)

    f = client.get("/api/insights/forecast", params={
        "outlet_id": outlet_id, "month": ym}).json()
    if f["so_far_rupees"] != dim * 1000:
        sheet = client.get("/api/sales/sheet", params={
            "outlet_id": outlet_id,
            "start": ym + "-01", "end": hi_prev.isoformat()}).json()
        nonzero = [(r["business_date"], r["channel_kind"], r["amount_rupees"],
                    r["source"]) for r in sheet["rows"][:8]]
        print("DEBUG so_far:", f["so_far_rupees"], "sample:", nonzero)
        print("DEBUG outlets seen by /outlets:",
              [o["id"] for o in client.get("/api/outlets").json()])
    assert abs(f["so_far_rupees"] - dim * 1000) < 0.01
    assert f["base_rate_rupees"] == 1000.0
    assert f["projected_rupees"] == round(dim * 1000.0, 2)

    # target + pace
    s = client.put("/api/insights/target", json={
        "outlet_id": outlet_id, "month": ym, "amount_rupees": 50000})
    assert s.status_code == 200
    f2 = client.get("/api/insights/forecast", params={
        "outlet_id": outlet_id, "month": ym}).json()
    assert f2["target_rupees"] == 50000.0, f2
    assert f2["percent_of_target"] == round(f2["projected_rupees"] / 50000 * 100, 1), f2


def test_anomalies_flags_cash_variance(client, outlet_id):
    client.post("/api/auth/stepup", json={"password": "change-me-please"})
    today = _today()
    cd = client.get("/api/cash/day", params={"outlet_id": outlet_id,
                                             "date": today}).json()
    cl = client.post("/api/cash/close", json={
        "outlet_id": outlet_id, "date": today,
        "counted_rupees": cd["expected_paise"] / 100 - 900,
        "note": "test leak"})
    if cl.status_code == 200:
        an = client.get("/api/insights/anomalies",
                        params={"outlet_id": outlet_id}).json()
        types = [a["type"] for a in an]
        assert "cash_variance" in types


def test_unit_economics_and_vendor_comparison(client, outlet_id):
    client.post("/api/auth/stepup", json={"password": "change-me-please"})
    cats = client.get("/api/lists/categories").json()
    raw_cat = next(c for c in cats if c["name"].lower().startswith("raw"))
    v1 = client.post("/api/vendors", json={"name": "UE Vendor A"}).json()
    v2 = client.post("/api/vendors", json={"name": "UE Vendor B"}).json()

    def buy(vendor_id, day, qty, total):
        return client.post("/api/expenses", json={
            "outlet_id": outlet_id, "business_date": day,
            "category_id": raw_cat["id"], "vendor_id": vendor_id,
            "amount_rupees": total, "mode": "cash",
            "item_name": "UE Basmati Rice", "quantity": qty, "unit": "kg"})

    r = buy(None, _today(), 20, 2000)          # no vendor + qty → rejected
    assert r.status_code == 422 and "vendor" in r.json()["detail"]
    buy(v1["id"], _old(6), 20, 2000)           # ₹100/kg
    buy(v1["id"], _today(), 25, 2500)          # ₹100/kg
    buy(v2["id"], _today(), 20, 1800)          # ₹90/kg ← cheapest

    ue = client.get("/api/insights/unit-economics", params={
        "start": _old(30), "end": _today(), "q": "UE Basmati"}).json()
    rice = next(x for x in ue if x["item"].lower() == "ue basmati rice")
    assert rice["purchases_count"] == 3
    assert rice["cheapest_vendor"] == "UE Vendor B"
    assert rice["cheapest_unit_price"] == 90.0
    assert rice["avg_unit_price"] == 96.67 or rice["avg_unit_price"] == 96.66


def test_benchmark_and_breakeven_shapes(client, outlet_id):
    b = client.get("/api/insights/benchmark", params={
        "start": _old(30), "end": _today()}).json()
    assert isinstance(b, list) and len(b) >= 1
    assert {"name", "sales_rupees", "best_weekday"} <= set(b[0].keys())

    be = client.get("/api/insights/breakeven", params={"outlet_id": outlet_id}).json()
    assert {"fixed_costs_rupees", "coverage_percent", "cost_per_bill_rupees"} \
        <= set(be.keys())


def test_budgets_flow(client):
    cats = client.get("/api/lists/categories").json()
    gas = next(c for c in cats if c["name"] == "Gas Cylinder")
    client.post("/api/auth/stepup", json={"password": "change-me-please"})
    r = client.put("/api/insights/budgets", json={
        "category_id": gas["id"], "monthly_rupees": 300})
    assert r.status_code == 200
    lst = client.get("/api/insights/budgets").json()
    row = next(x for x in lst if x["category"] == "Gas Cylinder")
    assert row["budget_rupees"] == 300.0
    assert isinstance(row["over"], bool)


def test_items_import_and_analytics(client, outlet_id):
    """dataio template roundtrip for item-level sales."""
    client.post("/api/auth/stepup", json={"password": "change-me-please"})
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Items"
    ws.append(["Date", "Item", "Category", "Qty", "Amount"])
    ws.append([_old(3), "Masala Papad", "Starters", 12, 1440])
    ws.append([_old(3), "Paneer Tikka", "Starters", 8, 2400])
    ws.append([_old(2), "Masala Papad", "Starters", 6, 720])
    x = _wb_bytes(wb)
    r = client.post(f"/api/data/import/sales_items?outlet_id={outlet_id}",
                    files={"file": ("items.xlsx", x, "application/octet-stream")})
    assert r.status_code == 200 and r.json()["created"] == 3

    ia = client.get("/api/insights/items", params={
        "start": _old(5), "end": _today(), "outlet_id": outlet_id}).json()
    top_names = [t["item"] for t in ia["top"]]
    assert top_names[0] == "Paneer Tikka"      # highest revenue first
    papad = next(t for t in ia["top"] if t["item"] == "Masala Papad")
    assert papad["qty"] == 18.0 and papad["sold_on_days"] == 2


def test_month_close_pack_has_close_sheets(client, outlet_id):
    client.post("/api/auth/stepup", json={"password": "change-me-please"})
    client.put("/api/sales/manual", json={
        "outlet_id": outlet_id, "business_date": _today(),
        "channel_kind": "other", "amount_rupees": 4321,
    })
    m = _today()[:7]
    r = client.get("/api/reports/ca-pack", params={
        "outlet_id": outlet_id, "month": m})
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    for sheet in ("Sales daily", "Bills", "Expenses", "Payroll",
                  "Vendor dues", "Advances", "Cash days"):
        assert sheet in wb.sheetnames, sheet
    sales_rows = list(wb["Sales daily"].iter_rows(min_row=2, values_only=True))
    manual = next(row for row in sales_rows
                  if row[0] == _today() and row[1] == "other"
                  and row[2] == "manual")
    assert manual[9] == 4321


def _old(days):
    from datetime import date, timedelta
    return (date.today() - timedelta(days=days)).isoformat()


def _wb_bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_remember_me_extends_token(client):
    r = client.post("/api/auth/login", json={
        "username": "owner", "password": "change-me-please", "remember": True})
    tok = r.json()["token"]
    import jwt as pyjwt
    from app.config import SECRET_KEY
    payload = pyjwt.decode(tok, SECRET_KEY, algorithms=["HS256"])
    import datetime as dt
    exp = dt.datetime.fromtimestamp(payload["exp"], tz=dt.timezone.utc)
    now = dt.datetime.now(dt.timezone.utc)
    hours = (exp - now).total_seconds() / 3600
    assert hours > 24 * 29, f"remember token should last ~30d, got {hours:.0f}h"
