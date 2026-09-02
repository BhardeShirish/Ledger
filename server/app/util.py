"""Business-date, money and time helpers. Timezone is pinned to IST."""
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

from fastapi import HTTPException
from zoneinfo import ZoneInfo

from .config import TZ_NAME

TZ = ZoneInfo(TZ_NAME)

DOW_MON_FIRST = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]

SPREADSHEET_HELP = (
    "That file could not be read as a spreadsheet. Upload an .xlsx file, or "
    "save your file as CSV and try again."
)


def read_sheet_rows(raw: bytes, *, sheet: str | None = None,
                    max_rows: int, max_cols: int,
                    label: str = "Workbook") -> list[list]:
    """Rows of cells from an uploaded spreadsheet, with a readable error.

    openpyxl raises zipfile.BadZipFile for anything that is not really an
    .xlsx, which reached the user as a bare "Internal server error". Every
    upload path goes through here so a wrong file type is always a clear 422.
    """
    import io

    from openpyxl import load_workbook

    if not raw:
        raise HTTPException(422, "The uploaded file is empty.")
    try:
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, SPREADSHEET_HELP)
    try:
        if not book.sheetnames:
            raise HTTPException(422, "That workbook has no sheets.")
        worksheet = (book[sheet] if sheet and sheet in book.sheetnames
                     else book[book.sheetnames[0]])
        if worksheet.max_row > max_rows or worksheet.max_column > max_cols:
            raise HTTPException(
                422, f"{label} exceeds the {max_rows:,}-row / "
                     f"{max_cols}-column limit")
        return [list(r) for r in worksheet.iter_rows(values_only=True)]
    finally:
        book.close()


def today_iso() -> str:
    return datetime.now(TZ).date().isoformat()


def now_local() -> datetime:
    return datetime.now(TZ)


def parse_date(s: str | None) -> date | None:
    if not s:
        return None
    try:
        return datetime.strptime(str(s)[:10], "%Y-%m-%d").date()
    except ValueError:
        return None


def month_bounds(year: int, month: int) -> tuple[date, date]:
    start = date(year, month, 1)
    end = date(year + (month == 12), (month % 12) + 1, 1) - timedelta(days=1)
    return start, end


def days_in_month(year: int, month: int) -> int:
    return month_bounds(year, month)[1].day


def rupees(paise: int) -> float:
    """Paise → decimal rupees for API responses."""
    return round(paise / 100.0, 2)


def paise(rupee_amount: float | int | str | None) -> int:
    """Rupees (from the UI) → integer paise, rounding half away from zero.

    Decimal keeps 0.145 → 15 rather than float's 14, and a non-finite amount
    is rejected here so it surfaces as a clear 422 instead of blowing up
    inside int() as a 500.
    """
    if rupee_amount is None or rupee_amount == "":
        return 0
    try:
        amount = Decimal(str(rupee_amount).strip())
    except (InvalidOperation, ValueError, TypeError):
        raise HTTPException(422, f"{rupee_amount!r} is not a valid amount")
    if not amount.is_finite():
        raise HTTPException(422, "Amount must be a finite number")
    return int((amount * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def hhmm_to_min(s: str | None) -> int | None:
    if not s:
        return None
    s = str(s).strip()
    try:
        if ":" in s:
            h, m = s.split(":")
            return (int(h) * 60 + int(m)) % 1440
        # compact form "822" → 8:22
        if len(s) >= 3:
            return (int(s[:-2]) * 60 + int(s[-2:])) % 1440
        return int(s) % 1440
    except ValueError:
        return None
