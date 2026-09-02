"""Monthly CA pack: XLSX with sales (tax split), expenses and payroll registers."""
from datetime import date
from io import BytesIO

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from sqlalchemy.orm import Session

from ..db import get_db
from ..models import (Advance, AdvanceRepayment, DayClosure, Employee, Expense,
                      ExpenseCategory, PayrollRun, Payslip, SalesBill,
                      SalesDaily, SalesItem, User, Vendor, VendorEntry)
from ..security import require_stepup
from ..util import month_bounds
from .helpers import assert_outlet_access

router = APIRouter(prefix="/reports", tags=["reports"])


def _append(sheet, values):
    sheet.append([
        "'" + value
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@"))
        else value
        for value in values
    ])


@router.get("/ca-pack")
def ca_pack(outlet_id: int, month: str, user: User = Depends(require_stepup),
            db: Session = Depends(get_db)):
    assert_outlet_access(db, user, outlet_id)
    y, m = map(int, month.split("-"))
    lo, hi = month_bounds(y, m)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sales daily"
    _append(ws, ["Date", "Channel", "Source", "Bills", "Gross", "Discount",
                 "Net", "Tax", "Tips", "Total"])
    rows = (db.query(SalesDaily)
              .filter_by(outlet_id=outlet_id)
              .filter(SalesDaily.business_date >= lo.isoformat(),
                      SalesDaily.business_date <= hi.isoformat())
              .order_by(SalesDaily.business_date).all())
    for r in rows:
        manual = r.source != "petpooja"
        gross = r.amount_paise if manual else r.gross_paise
        net = r.amount_paise if manual else r.net_paise
        total = r.amount_paise if manual else r.total_paise
        _append(ws, [r.business_date, r.channel_kind, r.source, r.bills,
                     gross / 100, r.discount_paise / 100, net / 100,
                     r.tax_paise / 100, r.tip_paise / 100, total / 100])

    ws2 = wb.create_sheet("Bills")
    heads = ["Invoice", "Timestamp", "Order type", "Area", "Persons",
             "Channel", "Gross", "Discount", "Net", "Tax", "Round off",
             "Tips", "Total"]
    _append(ws2, heads)
    bills = (db.query(SalesBill)
               .filter_by(outlet_id=outlet_id)
               .filter(SalesBill.business_date >= lo.isoformat(),
                       SalesBill.business_date <= hi.isoformat())
               .order_by(SalesBill.bill_ts).all())
    for b in bills:
        _append(ws2, [b.invoice_no, b.bill_ts, b.order_type, b.area, b.persons,
                      b.channel_kind, b.gross_paise / 100, b.discount_paise / 100,
                      b.net_paise / 100, b.tax_paise / 100, b.roundoff_paise / 100,
                      b.tip_paise / 100, b.total_paise / 100])

    ws3 = wb.create_sheet("Expenses")
    _append(ws3, ["Date", "Category", "Vendor", "Mode", "Description", "Amount"])
    cats = {c.id: c.name for c in db.query(ExpenseCategory).all()}
    exps = (db.query(Expense)
              .filter_by(outlet_id=outlet_id)
              .filter(Expense.business_date >= lo.isoformat(),
                      Expense.business_date <= hi.isoformat())
              .order_by(Expense.business_date).all())
    for e in exps:
        vendor = None
        if e.vendor_id:
            v = db.get(Vendor, e.vendor_id)
            vendor = v.name if v else None
        _append(ws3, [e.business_date, cats.get(e.category_id), vendor, e.mode,
                      e.description, round(e.amount_paise / 100, 2)])

    ws4 = wb.create_sheet("Payroll")
    _append(ws4, ["Employee", "Designation", "Salary", "÷Divisor", "Per day",
                  "Presents", "Halves", "Doubles", "Absents", "Credited days",
                  "Gross", "Bonus", "Deduction", "Advance recovery", "Net",
                  "Paid", "Mode"])
    run = (db.query(PayrollRun)
             .filter_by(outlet_id=outlet_id, year=y, month=m).first())
    if run:
        for s in db.query(Payslip).filter_by(run_id=run.id):
            emp = db.get(Employee, s.employee_id)
            per_day = s.monthly_salary_paise / (s.rate_divisor or 26) / 100
            _append(ws4, [
                emp.name if emp else "?", emp.designation if emp else "",
                round(s.monthly_salary_paise / 100, 2), s.rate_divisor,
                round(per_day, 2), s.presents, s.halves, s.doubles, s.absents,
                s.credited_days_x10 / 10.0, round(s.gross_paise / 100, 2),
                round(s.bonus_paise / 100, 2), round(s.deduction_paise / 100, 2),
                round(s.advance_recovery_paise / 100, 2), round(s.net_paise / 100, 2),
                round(s.paid_paise / 100, 2), s.mode])

    # Month-close extras
    bal_rows = []
    vendor_ids = {
        row.vendor_id
        for row in db.query(VendorEntry.vendor_id)
        .filter(VendorEntry.outlet_id == outlet_id,
                VendorEntry.date <= hi.isoformat())
    }
    for v in db.query(Vendor).filter(Vendor.id.in_(vendor_ids)).all():
        entries = (
            db.query(VendorEntry)
            .filter(VendorEntry.vendor_id == v.id,
                    VendorEntry.outlet_id == outlet_id,
                    VendorEntry.date <= hi.isoformat())
            .all()
        )
        bal = sum(x.amount_paise for x in entries if x.type == "purchase_credit") - \
              sum(x.amount_paise for x in entries if x.type != "purchase_credit")
        if bal:
            bal_rows.append([v.name, v.phone or "", round(bal / 100, 2)])
    ws5 = wb.create_sheet("Vendor dues")
    _append(ws5, ["Vendor", "Phone", "Balance due"])
    for r in sorted(bal_rows, key=lambda x: -x[2]):
        _append(ws5, r)

    ws6 = wb.create_sheet("Advances")
    _append(ws6, ["Date", "Employee", "Amount", "Remaining", "Status"])
    employees = {
        employee.id: employee
        for employee in db.query(Employee).filter_by(outlet_id=outlet_id).all()
    }
    advances = (
        db.query(Advance)
        .filter(Advance.employee_id.in_(employees), Advance.date <= hi.isoformat())
        .order_by(Advance.date.desc())
        .limit(300)
    )
    for a in advances:
        e = employees[a.employee_id]
        repaid = (
            db.query(AdvanceRepayment.amount_paise)
            .filter(AdvanceRepayment.advance_id == a.id,
                    AdvanceRepayment.date <= hi.isoformat())
            .all()
        )
        remaining = max(0, a.amount_paise - sum(row[0] for row in repaid))
        _append(ws6, [a.date, e.name if e else "?",
                      round(a.amount_paise / 100, 2),
                      round(remaining / 100, 2),
                      "cleared" if remaining == 0 else "open"])

    ws7 = wb.create_sheet("Cash days")
    _append(ws7, ["Date", "Expected", "Counted", "Variance", "Taken home",
                  "Left in drawer", "Note"])
    for c in (db.query(DayClosure)
                .filter_by(outlet_id=outlet_id)
                .filter(DayClosure.business_date >= lo.isoformat(),
                        DayClosure.business_date <= hi.isoformat())
                .order_by(DayClosure.business_date)):
        _append(ws7, [c.business_date, round(c.expected_cash_paise / 100, 2),
                      round(c.counted_cash_paise / 100, 2),
                      round(c.variance_paise / 100, 2),
                      round(c.moved_to_bank_paise / 100, 2),
                      round((c.counted_cash_paise - c.moved_to_bank_paise) / 100, 2),
                      c.note])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition":
                 f'attachment; filename="ca-pack-{month}-outlet{outlet_id}.xlsx"'})
