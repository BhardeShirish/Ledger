"""Excel data-in/out: export almost anything, download fillable templates,
and import them back — with the old-data rule enforced (only the owner may
bring in records older than the edit window)."""
import io
import hashlib

from fastapi import APIRouter, Depends, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from sqlalchemy import func
from sqlalchemy.orm import Session

from .. import bankstmt
from ..audit import audit, check_edit_window, get_setting_db
from ..db import get_db
from ..models import (Advance, Attendance, DayClosure, Employee, Expense,
                      ExpenseCategory, ImportBatch, PayrollRun, Payslip,
                      SalesBill, SalesDaily, SalesItem, Shift, User, Vendor,
                      VendorEntry)
from ..security import current_user
from ..util import SPREADSHEET_HELP, read_sheet_rows
from .helpers import assert_outlet_access, user_outlet_ids

router = APIRouter(prefix="/data", tags=["data"])


# ── helpers ───────────────────────────────────────────────────────────────

def _wb_response(wb: Workbook, filename: str) -> StreamingResponse:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'})


def _sheet(wb: Workbook, title: str, headers: list[str], rows) -> None:
    ws = wb.create_sheet(title)
    ws.append(headers)
    for r in rows:
        ws.append([_excel_safe(value) for value in r])


def _excel_safe(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _rupees(paise) -> float:
    return round((paise or 0) / 100, 2)


def _first_outlet(db: Session) -> int:
    o = db.query(Outlet_).filter_by(is_active=True).order_by(Outlet_.id).first()
    return o.id if o else 0


def _fmt_hhmm(m):
    if m is None:
        return ""
    m %= 1440
    return f"{m // 60:02d}:{m % 60:02d}"


from ..models import Outlet as Outlet_  # noqa: E402


# ── EXPORT definitions ────────────────────────────────────────────────────
# each: (sensitive, builder(db, user, params) -> (filename, [(title, headers, rows)]))

def _att_rows(db, params):
    lo, hi, oid = params["start"], params["end"], params["outlet_id"]
    emps = {e.id: e for e in db.query(Employee)
            .filter(Employee.outlet_id == oid)}
    rows = (db.query(Attendance)
              .filter(Attendance.business_date >= lo,
                      Attendance.business_date <= hi,
                      Attendance.employee_id.in_(emps.keys() or [0]))
              .order_by(Attendance.business_date, Attendance.employee_id).all())
    out = []
    for r in rows:
        e = emps.get(r.employee_id)
        out.append([r.business_date, e.code if e else "", e.name if e else "?",
                    r.status, _fmt_hhmm(r.in_min), _fmt_hhmm(r.out_min),
                    "Y" if r.double_duty else "", r.late_min, r.ot_min, r.note])
    return [("Attendance", ["Date", "Employee Code", "Employee Name", "Status",
                            "In (HH:MM)", "Out (HH:MM)", "Double Duty",
                            "Late Min", "OT Min", "Note"], out)]


EXPORTS = {}


def export_def(entity, sensitive=False):
    def deco(fn):
        EXPORTS[entity] = {"fn": fn, "sensitive": sensitive}
        return fn
    return deco


@export_def("analytics")
def exp_analytics(db, user, p):
    from .stats import analytics as analytics_fn
    data = analytics_fn(start=p["start"], end=p["end"],
                        outlet_id=p.get("outlet_id") or None,
                        user=user, db=db)
    keys = list(data["series"].keys())
    rows = []
    for i, day in enumerate(data["days"]):
        rows.append([day] + [data["series"][k][i] for k in keys])
    return f"analysis-{p['start']}_{p['end']}.xlsx", [(
        "Analysis", ["Date"] + keys,
        rows)]


@export_def("unit_economics")
def exp_unit_econ(db, user, p):
    from .insights import unit_economics
    data = unit_economics(start=p["start"], end=p["end"],
                          q=None, outlet_id=p.get("outlet_id") or None,
                          user=user, db=db)
    rows = []
    for it in data:
        for pu in it["purchases"]:
            rows.append([it["item"], it["category"], pu["date"], pu["vendor"],
                         pu["qty"], pu["unit"], pu["total_rupees"],
                         pu["unit_price_rupees"]])
    return "unit-economics.xlsx", [(
        "Unit prices",
        ["Item", "Category", "Date", "Vendor", "Qty", "Unit", "Total",
         "Per Unit"], rows)]


@export_def("items")
def exp_items(db, user, p):
    from .insights import item_analytics
    data = item_analytics(start=p["start"], end=p["end"],
                          outlet_id=p.get("outlet_id") or None,
                          user=user, db=db)
    rows = [[t["item"], t["category"], t["qty"], t["revenue_rupees"],
             t["sold_on_days"], t["menu_presence_percent"]]
            for t in data["top"]]
    for n in data["dead_items"]:
        rows.append([n, "", 0, 0, 0, 0])
    return "menu-items.xlsx", [(
        "Menu items", ["Item", "Category", "Qty", "Revenue", "Days Sold",
                       "Menu Presence %"], rows)]


@export_def("attendance")
def exp_attendance(db, user, p):
    assert_outlet_access(db, user, p["outlet_id"])
    return f"attendance-{p['start']}_{p['end']}.xlsx", _att_rows(
        db, {"start": p["start"], "end": p["end"], "outlet_id": p["outlet_id"]})


@export_def("employees", sensitive=True)
def exp_employees(db, user, p):
    oid = p.get("outlet_id") or (_first_outlet(db))
    assert_outlet_access(db, user, oid)
    backups = {e.id: e.name for e in db.query(Employee)}
    rows = []
    for e in db.query(Employee).filter_by(outlet_id=oid)\
                .order_by(Employee.code).all():
        rows.append([e.code, e.name, e.phone, e.designation,
                     _rupees(e.monthly_salary_paise), e.divisor,
                     e.join_date or "", e.working_status,
                     e.off_dow if e.off_dow is not None else "",
                     e.pref_off_dow if e.pref_off_dow is not None else "",
                     backups.get(e.backup_employee_id, ""), e.upi_id])
    return "employees.xlsx", [(
        "Staff", ["Code", "Name", "Phone", "Designation", "Monthly Salary",
                  "Divisor", "Join Date", "Status", "Off DOW (0=Mon)",
                  "Pref Off DOW", "Backup", "UPI ID"], rows)]


@export_def("expenses")
def exp_expenses(db, user, p):
    ids = user_outlet_ids(db, user)
    q = db.query(Expense).filter(Expense.outlet_id.in_(ids))
    if p.get("outlet_id"):
        assert_outlet_access(db, user, p["outlet_id"])
        q = q.filter(Expense.outlet_id == p["outlet_id"])
    if p.get("start"):
        q = q.filter(Expense.business_date >= p["start"])
    if p.get("end"):
        q = q.filter(Expense.business_date <= p["end"])
    cats = {c.id: c.name for c in db.query(ExpenseCategory)}
    vens = {v.id: v.name for v in db.query(Vendor)}
    rows = [[e.business_date, cats.get(e.category_id, ""),
             vens.get(e.vendor_id, "") if e.vendor_id else "",
             e.mode, _rupees(e.amount_paise), e.description,
             e.item_name, e.quantity, e.unit]
            for e in q.order_by(Expense.business_date).all()]
    return "expenses.xlsx", [("Expenses", ["Date", "Category", "Vendor", "Mode",
                                           "Amount", "Description", "Item",
                                           "Quantity", "Unit"], rows)]


@export_def("vendors")
def exp_vendors(db, user, p):
    rows = []
    for v in db.query(Vendor).order_by(Vendor.name).all():
        entries = db.query(VendorEntry).filter_by(vendor_id=v.id).all()
        bal = sum(a for t, a in ((e.type, e.amount_paise) for e in entries)
                  if t == "purchase_credit") - \
              sum(a for t, a in ((e.type, e.amount_paise) for e in entries)
                  if t != "purchase_credit")
        rows.append([v.name, v.phone, _rupees(bal),
                     "active" if v.is_active else "inactive"])
    return "vendors.xlsx", [("Vendors", ["Name", "Phone", "Balance Due",
                                         "Status"], rows)]


@export_def("sales_sheet")
def exp_sales_sheet(db, user, p):
    oid = p.get("outlet_id")
    assert_outlet_access(db, user, oid)
    q = db.query(SalesDaily).filter_by(outlet_id=oid)
    if p.get("start"):
        q = q.filter(SalesDaily.business_date >= p["start"])
    if p.get("end"):
        q = q.filter(SalesDaily.business_date <= p["end"])
    rows = [[r.business_date, r.channel_kind, r.source, r.bills,
             _rupees(r.gross_paise), _rupees(r.discount_paise),
             _rupees(r.net_paise), _rupees(r.tax_paise), _rupees(r.tip_paise),
             _rupees(r.total_paise if r.source == "petpooja" else r.amount_paise)]
            for r in q.order_by(SalesDaily.business_date,
                                SalesDaily.channel_kind).all()]
    return "sales.xlsx", [("Sales", ["Date", "Channel", "Source", "Bills",
                                     "Gross", "Discount", "Net", "Tax",
                                     "Tips", "Total"], rows)]


@export_def("bills")
def exp_bills(db, user, p):
    oid = p["outlet_id"]
    assert_outlet_access(db, user, oid)
    q = db.query(SalesBill).filter_by(outlet_id=oid)
    if p.get("date"):
        q = q.filter(SalesBill.business_date == p["date"])
    elif p.get("start"):
        q = q.filter(SalesBill.business_date >= p["start"],
                     SalesBill.business_date <= p.get("end"))
    rows = [[b.invoice_no, b.bill_ts, b.order_type, b.area, b.persons,
             b.channel_kind, _rupees(b.gross_paise), _rupees(b.discount_paise),
             _rupees(b.net_paise), _rupees(b.tax_paise), _rupees(b.tip_paise),
             _rupees(b.total_paise)]
            for b in q.order_by(SalesBill.bill_ts.desc()).limit(20000).all()]
    return "bills.xlsx", [("Bills", ["Invoice", "Timestamp", "Order Type",
                                     "Area", "Persons", "Channel", "Gross",
                                     "Discount", "Net", "Tax", "Tips",
                                     "Total"], rows)]


@export_def("closures")
def exp_closures(db, user, p):
    oid = p.get("outlet_id")
    assert_outlet_access(db, user, oid)
    rows = [[c.business_date, _rupees(c.expected_cash_paise),
             _rupees(c.counted_cash_paise), _rupees(c.variance_paise),
             _rupees(c.moved_to_bank_paise),
             _rupees(c.counted_cash_paise - c.moved_to_bank_paise),
             c.note] for c in db.query(DayClosure)
            .filter_by(outlet_id=oid)
            .order_by(DayClosure.business_date.desc()).limit(400).all()]
    return "cash-closures.xlsx", [("Cash Days", ["Date", "Expected", "Counted",
                                                 "Variance", "Taken Home",
                                                 "Left In Drawer", "Note"], rows)]


@export_def("payroll_run", sensitive=True)
def exp_payroll_run(db, user, p):
    run = db.get(PayrollRun, p["run_id"])
    if run is None:
        raise HTTPException(404, "Run not found")
    slips = db.query(Payslip).filter_by(run_id=run.id).all()
    emps = {e.id: e for e in db.query(Employee)}
    rows = []
    for s in slips:
        e = emps.get(s.employee_id)
        rows.append([e.name if e else "?", e.designation if e else "",
                     _rupees(s.monthly_salary_paise), s.rate_divisor,
                     s.presents, s.halves, s.doubles, s.absents,
                     s.lates_count, s.credited_days_x10 / 10.0,
                     _rupees(s.gross_paise), _rupees(s.bonus_paise),
                     _rupees(s.deduction_paise),
                     _rupees(s.advance_recovery_paise), _rupees(s.net_paise),
                     _rupees(s.paid_paise), s.mode or "", s.paid_on or ""])
    return f"payroll-{run.year}-{run.month:02d}.xlsx", [(
        f"Payroll {run.year}-{run.month:02d}",
        ["Name", "Designation", "Salary", "Divisor", "Presents", "Halves",
         "Doubles", "Absents", "Lates", "Credited Days", "Gross", "Bonus",
         "Deduction", "Advance Recovery", "Net", "Paid", "Mode", "Paid On"],
        rows)]


@export_def("advances", sensitive=True)
def exp_advances(db, user, p):
    emps = {e.id: e.name for e in db.query(Employee)}
    rows = [[a.date, emps.get(a.employee_id, "?"), _rupees(a.amount_paise),
             _rupees(a.remaining_paise), a.status, a.note]
            for a in db.query(Advance)
            .order_by(Advance.date.desc()).limit(500).all()]
    return "advances.xlsx", [("Advances", ["Date", "Employee", "Amount",
                                           "Remaining", "Status", "Note"], rows)]


@router.get("/export/{entity}.xlsx")
def do_export(entity: str,
              outlet_id: int | None = None,
              start: str | None = None,
              end: str | None = None,
              date: str | None = None,
              run_id: int | None = None,
              user: User = Depends(current_user),
              db: Session = Depends(get_db)):
    if entity not in EXPORTS:
        raise HTTPException(404, f"No export named '{entity}'")
    spec = EXPORTS[entity]
    if spec["sensitive"] and user.role != "owner":
        raise HTTPException(403, "Owner access required for this export")
    clean = {"outlet_id": outlet_id, "start": start, "end": end,
             "date": date, "run_id": run_id}
    clean = {k: v for k, v in clean.items() if v not in (None, "")}
    filename, sheets = spec["fn"](db, user, clean)
    wb = Workbook()
    wb.remove(wb.active)
    for title, headers, rows in sheets:
        _sheet(wb, title[:31], headers, rows)
    audit(db, None, user.id, "export", "data", entity)
    db.commit()
    return _wb_response(wb, filename)


# ── TEMPLATES ─────────────────────────────────────────────────────────────

IMPORT_ENTITIES = {
    "expenses": {
        "sheet": "Expenses",
        "headers": ["Date", "Category", "Vendor", "Mode", "Amount",
                    "Description", "Item", "Quantity", "Unit"],
        "example": [
            ["2000-01-01", "Gas Cylinder",
             "HP Gas Agency", "cash", 1150, "cylinder refill",
             "LPG cylinder", 1, "unit"],
            ["2000-01-01", "Vegetables & Fruits",
             "Sharma Mart", "credit", 3200, "weekly purchase"],
        ],
        "notes": [
            "Dates as YYYY-MM-DD.",
            "Modes: cash, upi, card, bank, credit, other.",
            "'credit' means bought on credit — it still counts as an expense "
            "immediately and raises the vendor's dues.",
            "New categories and vendors are created automatically from names.",
            "Rows dated older than the edit window need the OWNER password "
            "(upload while logged in as owner).",
        ],
    },
    "attendance": {
        "sheet": "Attendance",
        "headers": ["Date", "Employee Code", "Employee Name", "Status",
                    "In (HH:MM)", "Out (HH:MM)", "Double Duty", "Note"],
        "example": [
            ["2000-01-01", "2", "", "P", "07:05",
             "16:00", "", ""],
            ["2000-01-01", "", "Priya", "H", "07:00",
             "11:30", "Y", "left early"],
        ],
        "notes": [
            "Match staff by Code first, then by exact Name.",
            "Status: P, A, H (half), L (leave), WO (week off).",
            "Double Duty Y counts as one extra credited day.",
            "Blank In/Out is allowed only for A/L/WO.",
            "Older-than-window dates need the OWNER password.",
        ],
    },
    "sales_manual": {
        "sheet": "Sales",
        "headers": ["Date", "Cash", "UPI", "Card", "Other"],
        "example": [
            ["2000-01-01", 5000, 7000, 0, 0],
        ],
        "notes": [
            "One row per day; amounts in rupees.",
            "Fills your manual day-sheet numbers — Petpooja imports override "
            "these for days they cover.",
            "Older-than-window dates need the OWNER password.",
        ],
    },
    "sales_items": {
        "sheet": "Items",
        "headers": ["Date", "Item", "Category", "Qty", "Amount"],
        "example": [
            ["2000-01-01", "Masala Papad", "Starters", 4, 240],
        ],
        "notes": [
            "One row per item per day (Petpooja 'Item wise sales' export "
            "matches after deleting its extra columns).",
            "Powers top dishes, dead items and menu presence stats.",
            "Older-than-window dates need the OWNER password.",
        ],
    },
}


@router.get("/template/{entity}.xlsx")
def template(entity: str, user: User = Depends(current_user)):
    if entity not in IMPORT_ENTITIES:
        raise HTTPException(404, f"No template for '{entity}'")
    spec = IMPORT_ENTITIES[entity]
    wb = Workbook()
    info = wb.active
    info.title = "How to fill"
    info.append(["Follow these rules, then delete this sheet before uploading."])
    for n in spec["notes"]:
        info.append([f"• {n}"])
    _sheet(wb, spec["sheet"], spec["headers"], spec["example"])
    return _wb_response(wb, f"template-{entity}.xlsx")


# ── IMPORT ────────────────────────────────────────────────────────────────

def _read_grid(file_bytes: bytes, sheet_hint: str) -> list[list]:
    """Rows of cells from an uploaded sheet.

    Real .xlsx goes through openpyxl because only that honours the sheet name
    the downloadable template uses. Anything else - a CSV the user exported,
    an older .xls, an HTML table a bank named .xls - goes through the
    statement reader, which sniffs the actual format instead of trusting the
    extension. Without this, any non-xlsx upload died inside zipfile with a
    500 and no explanation.
    """
    if not file_bytes:
        raise HTTPException(422, "The uploaded file is empty.")
    if file_bytes[:4] != b"PK\x03\x04":
        try:
            return bankstmt.read_grid(file_bytes)
        except ValueError as ex:
            raise HTTPException(422, str(ex))
        except Exception:
            raise HTTPException(422, SPREADSHEET_HELP)
    return read_sheet_rows(file_bytes, sheet=sheet_hint,
                           max_rows=50_000, max_cols=100)


def _parse_rows(file_bytes: bytes, sheet_hint: str):
    rows = _read_grid(file_bytes, sheet_hint)
    if len(rows) > 50_000 or any(len(r) > 100 for r in rows):
        raise HTTPException(422,
                            "Workbook exceeds the 50,000-row / 100-column limit")
    hdr_i = next((i for i, r in enumerate(rows[:8])
                  if sum(1 for c in r if c not in (None, "")) >= 3), 0)
    headers = [str(c or "").strip().lower() for c in rows[hdr_i]] if rows else []
    data = []
    for r in rows[hdr_i + 1:]:
        if not any(c not in (None, "") for c in r):
            continue
        data.append({headers[j]: r[j] if j < len(r) else None
                     for j in range(len(headers))})
    return data


def _as_date(v):
    """Accept the date formats people actually have, not only ISO.

    A sheet exported from Excel in India says 05/01/2026, not 2026-01-05.
    Rejecting those produced "bad date" on every single row with no clue why.
    """
    iso = bankstmt.parse_date_cell(v)
    if iso is None:
        return None
    if int(iso[:4]) <= 2000:
        return None              # template example rows must be removed
    return iso


def _as_paise(v):
    """Money as paise, or None. Never raises.

    float("1,250.00") throws, and that reached the user as a 500.
    """
    return bankstmt.parse_amount_cell(v)


def _as_number(v):
    """A plain number (quantity), or None. Never raises."""
    if v in (None, ""):
        return None
    try:
        return float(str(v).replace(",", "").strip())
    except (TypeError, ValueError):
        return None


# A bank statement dropped on the generic importer used to fail every row
# with "bad date or amount", which explains nothing. These headers say the
# file belongs on the bank-statement page instead.
BANK_HEADER_HINTS = ("narration", "withdrawal", "deposit", "closing balance",
                     "particulars", "chq./ref", "value dt", "dr/cr")


def _looks_like_bank_statement(rows) -> bool:
    if not rows:
        return False
    keys = [str(k or "").lower() for k in rows[0]]
    hits = sum(1 for hint in BANK_HEADER_HINTS if any(hint in k for k in keys))
    return hits >= 2


def _old_data_guard(dates: list[str], user, db):
    cutoff_h = int(get_setting_db(db, "edit_cutoff_hours", 48))
    for d in sorted({x for x in dates if x}):
        try:
            check_edit_window(d, user, db)
        except HTTPException:
            raise HTTPException(
                403, f"Row dated {d} is older than the edit window. "
                     "Only the owner (after password verification) can import old data.")


@router.post("/import/{entity}")
async def run_import(entity: str, outlet_id: int, file: UploadFile,
                     user: User = Depends(current_user),
                     db: Session = Depends(get_db)):
    if entity not in IMPORT_ENTITIES:
        raise HTTPException(404, f"No importer for '{entity}'")
    assert_outlet_access(db, user, outlet_id)
    raw = await file.read(10 * 1024 * 1024 + 1)
    if len(raw) > 10 * 1024 * 1024:
        raise HTTPException(413, "Import file is larger than 10 MB")
    content_hash = hashlib.sha256(raw).hexdigest()
    duplicate = (
        db.query(ImportBatch)
        .filter_by(outlet_id=outlet_id, report_kind=f"generic:{entity}",
                   content_hash=content_hash, status="committed")
        .first()
    )
    if duplicate:
        return {"created": 0, "skipped": 0, "errors": [],
                "duplicate": True, "batch_id": duplicate.id}
    spec = IMPORT_ENTITIES[entity]
    rows = _parse_rows(raw, spec["sheet"])
    if not rows:
        raise HTTPException(422, "No data rows found under the header row")
    if _looks_like_bank_statement(rows):
        raise HTTPException(
            422,
            "This looks like a bank statement, not a Ledger sheet. Import it "
            "from Money → Bank statement — that page reads the narration, "
            "groups payments by who you paid, and remembers them for next time.")

    def col(r, *names):
        for n in names:
            for hk, hv in r.items():
                if hk.startswith(n.lower()) and hv not in (None, ""):
                    return hv
        return None

    created = skipped = 0
    errors = []

    if entity == "expenses":
        dates = [_as_date(col(r, "date")) for r in rows]
        _old_data_guard(dates, user, db)
        for i, r in enumerate(rows, start=2):
            d = _as_date(col(r, "date"))
            amt_paise = _as_paise(col(r, "amount"))
            if not d or not amt_paise:
                skipped += 1
                errors.append({
                    "row": i,
                    "why": "no valid date" if not d else "no valid amount",
                })
                continue
            cat_name = str(col(r, "category") or "Misc").strip()
            cat = db.query(ExpenseCategory).filter(
                func.lower(ExpenseCategory.name) == cat_name.lower()).first()
            if cat is None:
                cat = ExpenseCategory(name=cat_name)
                db.add(cat)
                db.flush()
            ven_name = str(col(r, "vendor") or "").strip()
            vid = None
            if ven_name:
                ven = db.query(Vendor).filter(
                    func.lower(Vendor.name) == ven_name.lower()).first()
                if ven is None:
                    ven = Vendor(name=ven_name)
                    db.add(ven)
                    db.flush()
                vid = ven.id
            mode = str(col(r, "mode") or "cash").strip().lower()
            if mode not in ("cash", "upi", "card", "bank", "credit", "other"):
                mode = "other"
            item_name = str(col(r, "item") or "").strip()
            quantity = _as_number(col(r, "quantity")) or None
            if quantity and (not vid or not item_name):
                skipped += 1
                errors.append({"row": i, "why": "quantity needs item and vendor"})
                continue
            expense = Expense(
                outlet_id=outlet_id, business_date=d,
                category_id=cat.id, vendor_id=vid,
                amount_paise=amt_paise, mode=mode,
                description=str(col(r, "description") or ""),
                item_name=item_name,
                quantity=quantity,
                unit=str(col(r, "unit") or "").strip(),
                entered_by=user.id,
            )
            db.add(expense)
            db.flush()
            if expense.quantity:
                from .inventory import ensure_purchase_movement
                ensure_purchase_movement(
                    db, outlet_id, d, expense.item_name, expense.quantity,
                    expense.unit, expense.amount_paise, user.id, expense.id,
                )
            if mode == "credit" and vid is not None:
                db.add(VendorEntry(
                    vendor_id=vid, outlet_id=outlet_id, date=d,
                    type="purchase_credit", amount_paise=expense.amount_paise,
                    expense_id=expense.id, note=expense.description,
                    entered_by=user.id,
                ))
            created += 1

    elif entity == "attendance":
        dates = [_as_date(col(r, "date")) for r in rows]
        _old_data_guard(dates, user, db)
        from datetime import date as _D

        from ..attendance_lib import resolve_shift
        emps_all = db.query(Employee).filter_by(outlet_id=outlet_id).all()
        by_code = {e.code.strip(): e for e in emps_all if e.code}
        by_name = {e.name.lower(): e for e in emps_all}
        for i, r in enumerate(rows, start=2):
            d = _as_date(col(r, "date"))
            code = str(col(r, "employee code") or "").strip()
            name = str(col(r, "employee name") or "").strip()
            emp = by_code.get(code) or by_name.get(name.lower())
            status = str(col(r, "status") or "P").strip().upper()[:2]
            if not d or emp is None or status not in ("P", "A", "H", "L", "WO"):
                skipped += 1
                errors.append({"row": i, "why": ("example/invalid date" if not d else f"unknown staff ({name or code})")})
                continue
            def hm(v):
                s = str(v or "").strip()
                if not s:
                    return None
                try:
                    hh, mm = s.split(":")
                    return int(hh) % 24 * 60 + int(mm) % 60
                except ValueError:
                    return None
            in_m, out_m = hm(col(r, "in")), hm(col(r, "out"))
            dd = str(col(r, "double duty") or "").strip().upper() == "Y"
            row = db.query(Attendance).filter_by(
                employee_id=emp.id, business_date=d).first()
            if row is None:
                row = Attendance(employee_id=emp.id, business_date=d)
                db.add(row)
            row.status = status
            sh = resolve_shift(db, emp, _D.fromisoformat(d))
            row.shift_id = None if sh in (None, "off") else sh.id
            row.in_min, row.out_min = in_m, out_m
            row.double_duty = dd and status in ("P", "H")
            eff_sh = sh if sh not in (None, "off") else None
            late = ot = 0
            if status in ("P", "H") and eff_sh and in_m is not None:
                late = max(0, in_m - (eff_sh.start_min + eff_sh.grace_min))
                if out_m is not None:
                    out_adj = out_m + 1440 if out_m < in_m else out_m
                    ot = max(0, out_adj - (eff_sh.end_min + eff_sh.ot_grace_min))
            row.late_min, row.ot_min = late, ot
            row.is_open = status in ("P", "H") and in_m is not None and out_m is None
            row.note = str(col(r, "note") or "")
            row.source = "manual"
            row.marked_by = user.id
            created += 1

    elif entity == "sales_manual":
        dates = [_as_date(col(r, "date")) for r in rows]
        _old_data_guard(dates, user, db)
        for i, r in enumerate(rows, start=2):
            d = _as_date(col(r, "date"))
            if not d:
                skipped += 1
                errors.append({"row": i, "why": "bad date"})
                continue
            any_amt = False
            for kind in ("cash", "upi", "card", "other"):
                paise = _as_paise(col(r, kind))
                if not paise:
                    continue
                any_amt = True
                row = (db.query(SalesDaily)
                         .filter_by(outlet_id=outlet_id, business_date=d,
                                    channel_kind=kind, source="manual").first())
                if row is None:
                    row = SalesDaily(outlet_id=outlet_id, business_date=d,
                                     channel_kind=kind, source="manual",
                                     amount_paise=0)
                    db.add(row)
                row.amount_paise = paise
            created += 1 if any_amt else 0
            skipped += 0 if any_amt else 1

    elif entity == "sales_items":
        dates = [_as_date(col(r, "date")) for r in rows]
        _old_data_guard(dates, user, db)
        for i, r in enumerate(rows, start=2):
            d = _as_date(col(r, "date"))
            name = str(col(r, "item") or "").strip()
            qty = _as_number(col(r, "qty", "quantity")) or 0
            paise = _as_paise(col(r, "amount")) or 0
            if not d or not name or (qty <= 0 and paise <= 0):
                skipped += 1
                errors.append({"row": i, "why": "example/invalid row"})
                continue
            row = (db.query(SalesItem)
                     .filter_by(outlet_id=outlet_id, business_date=d,
                                item_name=name).first())
            if row is None:
                row = SalesItem(outlet_id=outlet_id, business_date=d,
                                item_name=name, qty=0, amount_paise=0)
                db.add(row)
            row.category = str(col(r, "category") or "")
            row.qty = qty
            row.amount_paise = paise
            created += 1

    batch = ImportBatch(
        outlet_id=outlet_id, filename=file.filename or f"{entity}.xlsx",
        report_kind=f"generic:{entity}", uploaded_by=user.id,
        rows_total=len(rows), rows_ok=created, rows_skipped=skipped,
        content_hash=content_hash, status="committed",
    )
    db.add(batch)
    db.flush()
    audit(db, None, user.id, f"import-{entity}", "data", entity,
          after={"created": created, "skipped": skipped, "batch_id": batch.id})
    db.commit()
    return {"created": created, "skipped": skipped, "errors": errors,
            "duplicate": False, "batch_id": batch.id}
